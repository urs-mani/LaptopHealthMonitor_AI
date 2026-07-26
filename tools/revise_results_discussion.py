from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_FINAL.docx"
)
OUTPUT = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_FINAL_RESULTS_UPDATED.docx"
)


RESULT_PARAGRAPHS = [
    (
        "The completed prototype passed the defined project-verification checks and "
        "successfully constructed all twelve graphical-interface pages. During visible "
        "execution, the application acquired and displayed host-laptop information such "
        "as CPU utilisation, RAM utilisation, disk usage, battery level, available "
        "temperature, uptime, health trends, component risks, and preventive-maintenance "
        "recommendations. Storage devices, installed applications, system inventory, "
        "settings controls, and report-export functions were also demonstrated. The "
        "consolidated validation scenarios are recorded in Table 4.3, the verification "
        "proof is provided in Annexure-II as Figure A2.1, and the corresponding interface "
        "and operational-output evidence is provided in Screenshots A2.1 to A2.12."
    ),
    (
        "The battery-data pipeline processed 2,794 discharge-cycle records from 34 NASA "
        "batteries and produced 16 processed columns, including twelve model features and "
        "the required prediction targets. Complete battery identities were separated "
        "between training and testing: 26 batteries were used for training and 8 different "
        "batteries were held out for evaluation. This battery-group holdout is important "
        "because it prevents cycles from the same battery from appearing in both sets and "
        "reduces identity leakage. Dataset-validation and split evidence is shown in "
        "Screenshots A2.13(a) and A2.13(b), while the processed data and split identifiers "
        "are preserved in processed_cycle_dataset.csv and model_metrics.json."
    ),
    (
        "Three regression algorithms were evaluated under the same held-out split. Extra "
        "Trees produced the best result with MAE = 8.197 cycles, RMSE = 16.443 cycles, "
        "R2 = 0.886, and MAPE = 51.897%. By comparison, Gradient Boosting obtained "
        "RMSE = 25.160 and R2 = 0.733, while Random Forest obtained RMSE = 25.930 and "
        "R2 = 0.717. Therefore, Extra Trees was selected because it produced the lowest "
        "held-out RMSE and the highest R2. The MAE means that the selected model's "
        "prediction differed from the actual RUL by approximately 8.2 cycles on average. "
        "The R2 value indicates that about 88.6% of the held-out RUL variation was "
        "explained by the engineered features. The larger RMSE shows that a smaller number "
        "of records had greater errors, because squared error gives those deviations more "
        "weight. Full comparison evidence is provided in Screenshot A2.13(b), "
        "regression_metrics.csv, and test_predictions.csv."
    ),
    (
        "MAPE is higher than the other regression indicators because relative error grows "
        "rapidly when the true remaining-cycle value approaches zero near end of life. "
        "For this reason, MAPE was not used alone for model selection; MAE, RMSE, and R2 "
        "were interpreted together. This prevents a misleading conclusion based on one "
        "percentage metric and provides a more balanced assessment of prediction quality. "
        "The recorded metric values can be verified through Screenshot A2.13(b), "
        "regression_metrics.csv, model_metrics.json, and training_summary.json."
    ),
    (
        "The Random Forest health classifier achieved 94.29% accuracy, weighted precision "
        "of 0.945, weighted recall of 0.943, and weighted F1-score of 0.943 across the "
        "Critical, Healthy, and Warning classes. The confusion matrix contains 594 correct "
        "predictions out of 630 held-out records: 129 Critical, 292 Healthy, and 173 "
        "Warning records were classified correctly. The remaining errors comprised 3 "
        "Critical records classified as Healthy; 10 Healthy records classified as "
        "Critical and 7 as Warning; and 7 Warning records classified as Critical and 9 as "
        "Healthy. These results indicate strong aggregate class separation while also "
        "showing that the classifier is not error-free. Classification metrics and the "
        "confusion matrix are evidenced in Screenshots A2.13(b) and A2.13(c), "
        "model_metrics.json, and training_summary.json."
    ),
    (
        "The training workflow saved the selected regressor, health classifier, processed "
        "dataset, feature-importance values, held-out predictions, comparison metrics, and "
        "training summary. The application also generated monitoring and system-report "
        "outputs in the supported formats. These persisted artifacts make the result "
        "reviewable and reproducible instead of presenting only a final accuracy value. "
        "The training-completion and artifact-location evidence is shown in Screenshot "
        "A2.13(c), and the report interface is shown in Screenshot A2.10. The results "
        "demonstrate that the academic prototype and research pipeline operate correctly "
        "for the evaluated environment and NASA dataset; however, they do not constitute "
        "manufacturer-certified failure predictions for every commercial laptop battery."
    ),
]


SUMMARY_ROWS = [
    (
        "Application and GUI validation",
        "Imports passed; twelve pages constructed; live monitoring and utility interfaces operated.",
        "Confirms successful integration of monitoring, navigation, storage, software inventory, settings, and reports.",
        "Table 4.3; Figure A2.1; Screenshots A2.1-A2.12",
    ),
    (
        "Dataset processing",
        "2,794 cycle records from 34 batteries; 16 processed columns.",
        "Confirms successful discovery, parsing, cleaning, feature construction, and target generation.",
        "Screenshot A2.13(a); processed_cycle_dataset.csv",
    ),
    (
        "Grouped validation",
        "26 training batteries and 8 different held-out test batteries.",
        "Battery-level separation reduces identity leakage and tests generalisation to unseen battery groups.",
        "Screenshots A2.13(a-b); model_metrics.json",
    ),
    (
        "RUL regression",
        "Extra Trees: MAE 8.197, RMSE 16.443, R2 0.886, MAPE 51.897%.",
        "Extra Trees outperformed Random Forest and Gradient Boosting on held-out RMSE and R2.",
        "Screenshot A2.13(b); regression_metrics.csv; test_predictions.csv",
    ),
    (
        "Health classification",
        "Accuracy 94.29%; weighted F1-score 0.943; 594 of 630 records correct.",
        "Strong aggregate classification with limited confusion among Critical, Healthy, and Warning classes.",
        "Screenshots A2.13(b-c); model_metrics.json",
    ),
    (
        "Saved outputs and reports",
        "Models, metrics, predictions, feature importance, summaries, and system reports generated.",
        "Persisted artifacts provide reproducible evidence beyond the displayed interface values.",
        "Screenshot A2.10; Screenshot A2.13(c); ai/models; ai/reports",
    ),
]


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip() == exact_text]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph matching {exact_text!r}, found {len(matches)}"
        )
    return matches[0]


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            paragraph._p.remove(child)
    paragraph.add_run(text)


def insert_paragraph_before(reference: Paragraph, text: str, style_name: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    paragraph.style = style_name
    paragraph.add_run(text)
    return paragraph


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r_pr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def build_summary_table(document: Document, old_table) -> None:
    headers = [
        "Result Area",
        "Obtained Result",
        "Discussion / Interpretation",
        "Evidence Reference",
    ]
    widths = [1.20, 1.62, 2.68, 1.68]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = old_table.alignment

    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shade)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        for run in paragraph.runs:
            set_font(run, 8.4, bold=True)
    repeat_header(table.rows[0])

    for record in SUMMARY_ROWS:
        row = table.add_row()
        for index, text in enumerate(record):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            for run in paragraph.runs:
                set_font(run, 8.1, bold=(index == 0))

    total_twips = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "90")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    old_xml = old_table._tbl
    old_xml.addprevious(table._tbl)
    old_xml.getparent().remove(old_xml)


def main() -> None:
    document = Document(str(SOURCE))
    heading = find_paragraph(document, "4.8 Results and Discussion")
    caption = find_paragraph(document, "Table 4.4 Results and discussion summary")

    existing_body: list[Paragraph] = []
    current = heading._p.getnext()
    while current is not None and current is not caption._p:
        if current.tag == qn("w:p"):
            paragraph = Paragraph(current, heading._parent)
            if paragraph.text.strip():
                existing_body.append(paragraph)
        current = current.getnext()
    if len(existing_body) != 4:
        raise RuntimeError(
            f"Expected four existing Results and Discussion paragraphs, found {len(existing_body)}"
        )

    for paragraph, text in zip(existing_body, RESULT_PARAGRAPHS[:4]):
        replace_paragraph_text(paragraph, text)
    for text in RESULT_PARAGRAPHS[4:]:
        insert_paragraph_before(caption, text, "I3 Normal Body")

    replace_paragraph_text(
        caption, "Table 4.4 Results, Interpretation, and Evidence Summary"
    )
    caption.paragraph_format.keep_with_next = True

    old_summary = None
    for table in document.tables:
        if (
            table.rows
            and table.rows[0].cells
            and table.rows[0].cells[0].text.strip() == "Component / Output"
        ):
            old_summary = table
            break
    if old_summary is None:
        raise RuntimeError("Could not locate the existing Table 4.4")
    build_summary_table(document, old_summary)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
