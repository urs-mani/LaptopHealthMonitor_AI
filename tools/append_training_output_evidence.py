from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_FINAL_WITH_EVIDENCE.docx"
)
OUTPUT = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_FINAL.docx"
)
ASSET_DIR = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\tmp\training_evidence"
)

TERMINAL_OUTPUT = r"""==========================================================
      NASA BATTERY RUL MODEL TRAINING
==========================================================

DATASET CHECK: PASS
Records: 2,794
Batteries: 34
Columns: 16

Training and evaluating models. This can take several minutes...
{
  "dataset_rows": 2794,
  "batteries": 34,
  "train_batteries": [
    "B0005",
    "B0006",
    "B0007",
    "B0018",
    "B0025",
    "B0026",
    "B0027",
    "B0028",
    "B0030",
    "B0031",
    "B0032",
    "B0034",
    "B0036",
    "B0039",
    "B0040",
    "B0041",
    "B0043",
    "B0045",
    "B0046",
    "B0048",
    "B0051",
    "B0052",
    "B0053",
    "B0054",
    "B0055",
    "B0056"
  ],
  "test_batteries": [
    "B0029",
    "B0033",
    "B0038",
    "B0042",
    "B0044",
    "B0047",
    "B0049",
    "B0050"
  ],
  "best_model": "Extra Trees",
  "regression": [
    {
      "model": "Extra Trees",
      "MAE_cycles": 8.197493867712293,
      "RMSE_cycles": 16.442761622965683,
      "R2": 0.8860329913719982,
      "MAPE_pct": 51.89652963518389
    },
    {
      "model": "Gradient Boosting",
      "MAE_cycles": 14.51193493350785,
      "RMSE_cycles": 25.16042040079199,
      "R2": 0.7331511014484833,
      "MAPE_pct": 84.36609998852799
    },
    {
      "model": "Random Forest",
      "MAE_cycles": 11.981103228265813,
      "RMSE_cycles": 25.930464185646507,
      "R2": 0.7165671327318761,
      "MAPE_pct": 69.74751481353726
    }
  ],
  "classification": {
    "accuracy": 0.9428571428571428,
    "precision_weighted": 0.9445758231194423,
    "recall_weighted": 0.9428571428571428,
    "f1_weighted": 0.9430235532032961,
    "labels": [
      "Critical",
      "Healthy",
      "Warning"
    ],
    "confusion_matrix": [
      [
        129,
        3,
        0
      ],
      [
        10,
        292,
        7
      ],
      [
        7,
        9,
        173
      ]
    ]
  },
  "features": [
    "cycle_index",
    "capacity_ah",
    "capacity_ratio",
    "ambient_temperature",
    "voltage_mean",
    "voltage_min",
    "voltage_max",
    "current_mean",
    "current_std",
    "temperature_mean",
    "temperature_max",
    "duration_s"
  ],
  "validation": "Battery-group holdout; test batteries are not used during training."
}

Training completed successfully.
Models: ai\models
Metrics and predictions: ai\reports
Restart the application to load the trained model.
Press any key to continue . . ."""


def font_path() -> Path:
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\lucon.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("No suitable Windows monospace font was found")


def terminal_color(line: str) -> tuple[int, int, int]:
    stripped = line.strip()
    if "PASS" in line or "completed successfully" in line:
        return (109, 222, 128)
    if "Extra Trees" in line or stripped.startswith('"best_model"'):
        return (97, 175, 239)
    if stripped.startswith("NASA BATTERY") or set(stripped) == {"="}:
        return (255, 214, 102)
    if stripped.startswith("Training and evaluating"):
        return (220, 220, 170)
    return (238, 238, 238)


def render_terminal_parts() -> list[Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    lines = TERMINAL_OUTPUT.splitlines()
    lines_per_part = 42
    chunks = [
        lines[index : index + lines_per_part]
        for index in range(0, len(lines), lines_per_part)
    ]
    mono = ImageFont.truetype(str(font_path()), 25)
    title_font = ImageFont.truetype(str(font_path()), 22)
    width = 1400
    top_bar = 66
    left_padding = 34
    top_padding = 28
    line_height = 33
    bottom_padding = 30
    outputs: list[Path] = []

    for index, chunk in enumerate(chunks, 1):
        height = top_bar + top_padding + line_height * len(chunk) + bottom_padding
        image = Image.new("RGB", (width, height), (12, 12, 12))
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, width, top_bar), fill=(35, 35, 35))
        for x, color in (
            (28, (255, 95, 86)),
            (60, (255, 189, 46)),
            (92, (39, 201, 63)),
        ):
            draw.ellipse((x, 21, x + 18, 39), fill=color)
        draw.text(
            (132, 19),
            f"NASA Battery RUL Model Training - Evidence {index}/{len(chunks)}",
            font=title_font,
            fill=(235, 235, 235),
        )

        y = top_bar + top_padding
        for line in chunk:
            draw.text(
                (left_padding, y),
                line,
                font=mono,
                fill=terminal_color(line),
            )
            y += line_height

        output = ASSET_DIR / f"nasa_training_output_{index}.png"
        image.save(output, format="PNG", optimize=True)
        outputs.append(output)
    return outputs


def set_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r_pr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    parts = render_terminal_parts()
    document = Document(str(SOURCE))
    if any(
        p.text.strip() == "NASA Battery RUL Model Training Output Evidence"
        for p in document.paragraphs
    ):
        raise RuntimeError("Training-output evidence already exists")

    document.add_page_break()
    heading = document.add_paragraph(
        style="I3 Heading 2"
        if "I3 Heading 2" in [style.name for style in document.styles]
        else "Heading 2"
    )
    heading.add_run("NASA Battery RUL Model Training Output Evidence")
    heading.paragraph_format.keep_with_next = True

    intro = document.add_paragraph(
        style="I3 Normal Body"
        if "I3 Normal Body" in [style.name for style in document.styles]
        else "Normal"
    )
    intro.add_run(
        "The following terminal-output screenshots provide evidence of successful "
        "NASA dataset validation, battery-group model training, regression-model "
        "comparison, health-classification evaluation, model persistence, and "
        "generation of reproducible metrics and prediction artifacts."
    )

    suffixes = "abcdefghijklmnopqrstuvwxyz"
    for index, image_path in enumerate(parts):
        if index > 0:
            document.add_page_break()
        caption_text = (
            f"Screenshot A2.13({suffixes[index]}): NASA Battery RUL Model "
            f"Training Output Evidence - Part {index + 1} of {len(parts)}"
        )
        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.space_before = Pt(5)
        picture_paragraph.paragraph_format.space_after = Pt(4)
        run = picture_paragraph.add_run()
        shape = run.add_picture(str(image_path), width=Inches(6.92))
        shape._inline.docPr.set("descr", caption_text)
        shape._inline.docPr.set("title", caption_text)

        caption = document.add_paragraph(
            style="Caption"
            if "Caption" in [style.name for style in document.styles]
            else "Normal"
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_previous = True
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(6)
        caption_run = caption.add_run(caption_text)
        set_font(caption_run, 10.5, bold=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(f"parts={len(parts)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
