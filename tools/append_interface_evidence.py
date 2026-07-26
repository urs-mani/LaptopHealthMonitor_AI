from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_MODIFIED.docx"
)
OUTPUT = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_FINAL_WITH_EVIDENCE.docx"
)

SCREENSHOTS = [
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-a350538b-fb1f-4b85-a4aa-282b9d470221.png"
        ),
        "Screenshot A2.1: Dashboard and Live System Health Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-045cb8e5-afbd-4d51-8ebc-acf5c048053e.png"
        ),
        "Screenshot A2.2: Real-Time System Monitoring Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-00c080b1-0355-4c99-9274-ffc8a6954f19.png"
        ),
        "Screenshot A2.3: Temperature Monitoring Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-da2f44fb-1793-4375-bc3c-409e3ad2c482.png"
        ),
        "Screenshot A2.4: RUL Prediction and AI Model Status Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-14af47ba-30d8-4bca-8fec-df2bd6c51f21.png"
        ),
        "Screenshot A2.5: System Health History Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-1e872477-92ad-493c-91cc-ab45c48baba1.png"
        ),
        "Screenshot A2.6: Preventive-Maintenance Recommendations Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-18f4bd86-b0cc-4147-b4ec-fbeb98768b2a.png"
        ),
        "Screenshot A2.7: Storage Manager Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-1f237a6b-ceda-4d24-a295-01a3aacb1415.png"
        ),
        "Screenshot A2.8: Smart Uninstaller and Installed-Program Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-fba5bd80-5a9a-4c7a-a2e0-7f31c24ef79c.png"
        ),
        "Screenshot A2.9: Complete System Data Output",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-209ccc14-3326-4664-b794-7a04a8d4c9ea.png"
        ),
        "Screenshot A2.10: Report Generation and Export Interface",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-11a7e0c7-4bbb-4ca2-8925-858963ab35c9.png"
        ),
        "Screenshot A2.11: Application Settings Interface",
    ),
    (
        Path(
            r"C:\Users\chand\AppData\Local\Temp"
            r"\codex-clipboard-25cacaa6-4a99-46c8-9665-ffee949ad25c.png"
        ),
        "Screenshot A2.12: About and System Information Interface",
    ),
]


def set_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r_pr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def main() -> None:
    missing = [str(path) for path, _ in SCREENSHOTS if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing screenshots:\n" + "\n".join(missing))

    document = Document(str(SOURCE))
    if any(
        p.text.strip() == "Annexure-II(B): Application Interface and Output Evidence"
        for p in document.paragraphs
    ):
        raise RuntimeError("Evidence subsection already exists in the source document")

    document.add_page_break()
    heading = document.add_paragraph(
        style="I3 Heading 2"
        if "I3 Heading 2" in [style.name for style in document.styles]
        else "Heading 2"
    )
    heading.add_run("Annexure-II(B): Application Interface and Output Evidence")
    heading.paragraph_format.keep_with_next = True

    intro = document.add_paragraph(
        style="I3 Normal Body"
        if "I3 Normal Body" in [style.name for style in document.styles]
        else "Normal"
    )
    intro.add_run(
        "The following screenshots provide evidence of the actual interfaces and "
        "outputs obtained from the implemented Laptop Health Monitor application. "
        "They cover real-time monitoring, temperature analysis, Remaining Useful "
        "Life assessment, historical trends, recommendations, storage and software "
        "inventory, system data, report generation, configuration, and application "
        "information."
    )

    for index, (image_path, caption_text) in enumerate(SCREENSHOTS):
        if index > 0:
            document.add_page_break()

        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.space_before = Pt(6)
        picture_paragraph.paragraph_format.space_after = Pt(5)
        run = picture_paragraph.add_run()
        inline_shape = run.add_picture(str(image_path), width=Inches(6.92))
        inline_shape._inline.docPr.set("descr", caption_text)
        inline_shape._inline.docPr.set("title", caption_text)

        caption = document.add_paragraph(
            style="Caption"
            if "Caption" in [style.name for style in document.styles]
            else "Normal"
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(6)
        caption.paragraph_format.keep_with_previous = True
        caption_run = caption.add_run(caption_text)
        set_font(caption_run, 11, bold=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
