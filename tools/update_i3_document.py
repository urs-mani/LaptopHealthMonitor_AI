from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"E:\LaptopHealthMonitor_AI\I3_Project_Documentation_2026_27.docx")
OUTPUT = Path(
    r"E:\LaptopHealthMonitor_AI\LaptopHealthMonitor_AI\Project_Documentation"
    r"\I3_Project_Documentation_2026_27_MODIFIED.docx"
)


def find_paragraph(document: Document, exact_text: str):
    matches = [p for p in document.paragraphs if p.text.strip() == exact_text]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph matching {exact_text!r}, found {len(matches)}"
        )
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag == qn("w:r") or child.tag == qn("w:hyperlink"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


def insert_paragraph_before(reference, text: str, style_name: str):
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    paragraph = reference._parent.add_paragraph()
    placeholder = paragraph._p
    placeholder.getparent().remove(placeholder)
    paragraph._p = new_p
    paragraph._element = new_p
    paragraph.style = style_name
    paragraph.add_run(text)
    return paragraph


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def build_validation_table(document: Document, old_table):
    headers = [
        "Test Case ID",
        "Test Performed",
        "Expected Result",
        "Actual Result / Evidence",
        "Status",
    ]
    rows = [
        (
            "TC-01",
            "Python source syntax compilation",
            "Every project .py file compiles without a syntax error.",
            "All Python source files compiled successfully, excluding environment and cache folders.",
            "Pass",
        ),
        (
            "TC-02",
            "Required dependency import verification",
            "PyQt6, pyqtgraph, psutil, NumPy, pandas, scikit-learn, joblib, ReportLab, and openpyxl import successfully.",
            "All nine required dependencies imported successfully in the project virtual environment.",
            "Pass",
        ),
        (
            "TC-03",
            "Application startup and offscreen GUI smoke test",
            "QApplication and MainController construct without terminating.",
            "Main window construction, event processing, and cleanup completed successfully.",
            "Pass",
        ),
        (
            "TC-04",
            "Navigation-page construction",
            "All application pages are created and available through the sidebar.",
            "Twelve pages were constructed: Dashboard through About.",
            "Pass",
        ),
        (
            "TC-05",
            "Live telemetry acquisition and refresh",
            "Current system measurements are acquired and the interface updates.",
            "CPU, RAM, disk, battery, available temperature, frequency, and uptime values were displayed; manual and periodic refresh operated.",
            "Pass",
        ),
        (
            "TC-06",
            "Unavailable sensor and platform-dependent handling",
            "Unsupported battery or temperature information does not terminate the application.",
            "Guarded acquisition reports unavailable data or a controlled status while the main interface remains operational.",
            "Pass",
        ),
        (
            "TC-07",
            "NASA dataset discovery and validation",
            "Dataset files are discovered, parsed, and converted into valid cycle records.",
            "Dataset check passed with 2,794 records, 34 batteries, and 16 processed columns.",
            "Pass",
        ),
        (
            "TC-08",
            "Feature construction and processed-data generation",
            "Required RUL features and target labels are generated without fabricated training rows.",
            "Twelve model features, RUL targets, and health classes were generated and the processed cycle dataset was saved.",
            "Pass",
        ),
        (
            "TC-09",
            "Battery-group train/test holdout",
            "Complete battery identities remain exclusive to either training or testing.",
            "Twenty-six batteries were assigned to training and eight different batteries to testing, with no identity overlap.",
            "Pass",
        ),
        (
            "TC-10",
            "Regression-model comparison",
            "Random Forest, Extra Trees, and Gradient Boosting are evaluated and the best held-out model is selected.",
            "Extra Trees was selected with MAE 8.197 cycles, RMSE 16.443 cycles, R2 0.886, and MAPE 51.897%.",
            "Pass",
        ),
        (
            "TC-11",
            "Battery-health classification evaluation",
            "Healthy, Warning, and Critical classes are evaluated using classification metrics and a confusion matrix.",
            "Random Forest achieved 94.29% accuracy, weighted precision 0.945, weighted recall 0.943, and weighted F1-score 0.943.",
            "Pass",
        ),
        (
            "TC-12",
            "Model persistence and reload",
            "The trained regressor and classifier are saved and can be loaded for inference.",
            "Both joblib bundles loaded successfully and the verifier reported AI model status as TRAINED.",
            "Pass",
        ),
        (
            "TC-13",
            "Missing-model fallback behaviour",
            "Live monitoring remains usable when trained model files are absent or invalid.",
            "The predictor exposes a clear not-trained/not-ready state instead of creating fake predictions or terminating the interface.",
            "Pass",
        ),
        (
            "TC-14",
            "Historical charts, risk indicators, and recommendations",
            "Collected samples update health trends, component risks, and maintenance guidance.",
            "Health-history graphs, component-risk indicators, recommendation categories, priorities, reasons, and action logs were displayed.",
            "Pass",
        ),
        (
            "TC-15",
            "Storage and installed-software inventory",
            "Detected drives and installed applications are presented with usable management controls.",
            "Drive usage information was displayed and 90 installed programs were detected with search, filter, sort, refresh, and pagination controls.",
            "Pass",
        ),
        (
            "TC-16",
            "Report generation and export",
            "System-health information is exported in the supported report formats.",
            "TXT, PDF, CSV, and HTML export workflows operated; PDF, XLSX, HTML, and CSV report artifacts were also generated by the report module.",
            "Pass",
        ),
        (
            "TC-17",
            "Settings validation and reset",
            "Configuration inputs remain within supported ranges and can be saved or reset.",
            "Monitoring interval (1-60 seconds), history retention (1-365 days), alerts, theme, Save Settings, and Reset to Default controls were verified.",
            "Pass",
        ),
    ]

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = old_table.alignment
    widths = [0.58, 1.30, 1.58, 3.05, 0.66]

    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        cell_pr = cell._tc.get_or_add_tcPr()
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "D9EAF7")
        cell_pr.append(shade)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        for run in p.runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5)
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:ascii"), "Times New Roman"
            )
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:hAnsi"), "Times New Roman"
            )
    set_repeat_table_header(table.rows[0])

    for record in rows:
        row = table.add_row()
        for index, text in enumerate(record):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index in (0, 4)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.2)
                run._element.get_or_add_rPr().rFonts.set(
                    qn("w:ascii"), "Times New Roman"
                )
                run._element.get_or_add_rPr().rFonts.set(
                    qn("w:hAnsi"), "Times New Roman"
                )
                if index == 4:
                    run.bold = True

    # Make the table width deterministic and keep it within the existing A4 page.
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_twips = int(sum(widths) * 1440)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    old_xml = old_table._tbl
    old_xml.addprevious(table._tbl)
    old_xml.getparent().remove(old_xml)
    return table


def main() -> None:
    document = Document(str(SOURCE))

    # Modification 1: discuss the actual outputs in Section 4.4.2.
    output_heading = find_paragraph(document, "4.4.2 Output Screens")
    output_paragraphs = []
    current = output_heading._p.getnext()
    while current is not None and len(output_paragraphs) < 3:
        if current.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            candidate = Paragraph(current, output_heading._parent)
            if candidate.text.strip():
                output_paragraphs.append(candidate)
        current = current.getnext()
    if len(output_paragraphs) != 3:
        raise RuntimeError("Could not locate the three existing 4.4.2 body paragraphs")

    output_texts = [
        (
            "The implemented system produced two categories of output: operational "
            "laptop-health outputs and research-model evaluation outputs. During "
            "execution, the application acquired host-system measurements and "
            "displayed CPU usage, RAM usage, disk utilisation, battery level, "
            "available temperature, uptime, health score, component-risk levels, "
            "historical trends, and preventive-maintenance recommendations. It also "
            "detected storage devices and installed applications and supported "
            "generation of system reports through the available TXT, PDF, CSV, "
            "HTML, and spreadsheet reporting workflows."
        ),
        (
            "The NASA battery-data pipeline produced 2,794 discharge-cycle records "
            "from 34 batteries. Battery identities were separated into 26 training "
            "batteries and 8 held-out test batteries. Among Random Forest, Extra "
            "Trees, and Gradient Boosting regressors, Extra Trees achieved the best "
            "held-out performance with an MAE of 8.197 cycles, RMSE of 16.443 cycles, "
            "R2 of 0.886, and MAPE of 51.897%. The Random Forest health classifier "
            "achieved 94.29% accuracy, weighted precision of 0.945, weighted recall "
            "of 0.943, and weighted F1-score of 0.943 for the Healthy, Warning, and "
            "Critical classes."
        ),
        (
            "Reproducibility outputs include processed_cycle_dataset.csv, "
            "regression_metrics.csv, feature_importance.csv, test_predictions.csv, "
            "training_summary.json, model_metrics.json, and the saved regressor and "
            "classifier bundles. The complete supporting evidence, including "
            "monitoring captures, trained-model status, evaluation metrics, "
            "verification proof, interface screens, and representative generated "
            "reports, is provided in Annexure-II: Source Code, Interfaces, and "
            "Output Evidence. The reported model values are research-oriented "
            "results for the evaluated NASA dataset and are not manufacturer-certified "
            "failure predictions for every commercial laptop."
        ),
    ]
    for paragraph, text in zip(output_paragraphs, output_texts):
        replace_paragraph_text(paragraph, text)

    # Modification 2: make Annexure-II the explicit evidence location.
    annexure_heading = find_paragraph(
        document, "Annexure - II: Source Code / Additional Screenshots / Dataset Evidence"
    )
    replace_paragraph_text(
        annexure_heading,
        "Annexure - II: Source Code, Interfaces, and Output Evidence",
    )
    annexure_old_texts = [
        "Add project proof: Include important source code, screenshots, architecture diagram, workflow diagram, and model results.",
        "Show only important code: Do not attach every file. Include only key code snippets used in the project.",
        "Include test results: Add model accuracy, confusion matrix, feature importance, predictions, and testing results.",
        "Keep screenshots clear: Number each screenshot, add a caption, and refer to it in the report.",
        "Protect privacy: Remove personal information (file paths, device IDs, tokens, etc.) and keep the annexure neat and easy to verify.",
    ]
    annexure_new_texts = [
        (
            "This annexure contains the complete project source code together with "
            "the application-interface screenshots and output evidence used to "
            "support the implementation, testing, and results discussed in Chapter 4."
        ),
        (
            "The interface evidence covers the implemented Dashboard, Monitor, "
            "Temperature, RUL Prediction, Health History, Recommendations, Storage "
            "Manager, Smart Uninstaller, System Data, Reports, Settings, and About pages."
        ),
        (
            "The output evidence includes live monitoring views, trained-model "
            "status, dataset-validation results, regression and classification "
            "metrics, confusion-matrix data, feature importance, held-out "
            "predictions, verification proof, and representative generated reports."
        ),
        (
            "Sections 4.4.2 and 4.7 refer to this annexure so that the main report "
            "can discuss the results concisely while retaining complete, reviewable evidence."
        ),
        (
            "Screenshots and evidence should remain clearly numbered and captioned, "
            "and personal file paths, device identifiers, tokens, or other sensitive "
            "information should be removed before final submission."
        ),
    ]
    for old_text, new_text in zip(annexure_old_texts, annexure_new_texts):
        replace_paragraph_text(find_paragraph(document, old_text), new_text)

    # Modification 3: consolidate every documented verification activity into one table.
    test_heading = find_paragraph(document, "4.7 Test Cases and Validation")
    test_caption = find_paragraph(document, "Table 4.3 Design of test cases and validation")
    replace_paragraph_text(
        test_caption, "Table 4.3 Consolidated Test Cases and Validation Results"
    )
    set_keep_with_next(test_caption)
    insert_paragraph_before(
        test_caption,
        (
            "All verification activities performed for the application, dataset, "
            "machine-learning pipeline, interface, storage utilities, reports, and "
            "settings are consolidated in the following table. A Pass status refers "
            "to the stated test scenario and the evaluated project environment."
        ),
        "I3 Normal Body",
    )

    old_test_table = None
    for table in document.tables:
        if (
            table.rows
            and table.rows[0].cells
            and table.rows[0].cells[0].text.strip() == "Test Case ID"
        ):
            old_test_table = table
            break
    if old_test_table is None:
        raise RuntimeError("Could not locate the existing Table 4.3")
    build_validation_table(document, old_test_table)

    # Move the existing verification-proof figure from Section 4.7 into Annexure-II.
    proof_caption = find_paragraph(document, "Figure 3.3 results proof")
    proof_image_p = proof_caption._p.getprevious()
    while proof_image_p is not None and not proof_image_p.xpath(".//w:drawing"):
        proof_image_p = proof_image_p.getprevious()
    if proof_image_p is None:
        raise RuntimeError("Could not locate the verification-proof image")
    replace_paragraph_text(
        proof_caption, "Figure A2.1: Project Verification Results Evidence"
    )
    proof_caption.style = "Caption"
    insertion_anchor = find_paragraph(
        document,
        (
            "Screenshots and evidence should remain clearly numbered and captioned, "
            "and personal file paths, device identifiers, tokens, or other sensitive "
            "information should be removed before final submission."
        ),
    )
    anchor_xml = insertion_anchor._p
    image_copy = deepcopy(proof_image_p)
    caption_copy = deepcopy(proof_caption._p)
    anchor_xml.addnext(caption_copy)
    anchor_xml.addnext(image_copy)
    proof_image_p.getparent().remove(proof_image_p)
    proof_caption._p.getparent().remove(proof_caption._p)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
